"""Export a range of workouts to XLSX for external analysis.

One row per set with workout-level metadata repeated across the workout's rows
(easier for pivot tables and consistent with the import format). Includes
computed analysis columns (effective load, volume, e1RM) that respect the
bodyweight-aware rules used elsewhere in the app.
"""

import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.body_metric import BodyMetric
from app.models.workout import Set, Workout, WorkoutExercise
from app.services.pr_detection import effective_e1rm, effective_load


HEADERS = [
    "Date",
    "Workout Name",
    "Bodyweight (lbs)",
    "Sleep (1-5)",
    "Fatigue (1-5)",
    "Exercise",
    "Equipment",
    "Set",
    "Weight (lbs)",
    "Reps",
    "RPE",
    "Type",
    "Effective Load (lbs)",
    "Volume (lbs)",
    "e1RM (lbs)",
    "Is PR",
    "Workout Notes",
    "Exercise Notes",
    "Set Notes",
]


async def export_workouts_to_xlsx(
    db: AsyncSession,
    user_id: int,
    start_date: date | None,
    end_date: date | None,
) -> bytes:
    conditions = [Workout.user_id == user_id]
    if start_date:
        conditions.append(Workout.date >= start_date)
    if end_date:
        conditions.append(Workout.date <= end_date)

    stmt = (
        select(Workout)
        .where(and_(*conditions))
        .options(
            selectinload(Workout.exercises).selectinload(WorkoutExercise.sets),
            selectinload(Workout.exercises).selectinload(WorkoutExercise.exercise),
        )
        .order_by(Workout.date, Workout.id)
    )
    workouts_list = (await db.execute(stmt)).scalars().all()

    bm_rows = (await db.execute(
        select(BodyMetric.date, BodyMetric.bodyweight_lbs)
        .where(BodyMetric.user_id == user_id, BodyMetric.bodyweight_lbs.is_not(None))
        .order_by(BodyMetric.date)
    )).all()
    bm_list = [(r.date, float(r.bodyweight_lbs)) for r in bm_rows]

    def bw_for_date(d: date) -> float | None:
        result = None
        for bm_date, bw in bm_list:
            if bm_date <= d:
                result = bw
            else:
                break
        return result

    wb = Workbook()
    ws = wb.active
    ws.title = "Workouts"
    ws.append(HEADERS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for w in workouts_list:
        bw = float(w.bodyweight_lbs) if w.bodyweight_lbs else bw_for_date(w.date)
        for we in sorted(w.exercises, key=lambda x: x.order_index):
            equipment = we.exercise.equipment if we.exercise else None
            exercise_name = we.exercise.name if we.exercise else f"Exercise #{we.exercise_id}"
            for s in sorted(we.sets, key=lambda x: x.set_number):
                weight = float(s.weight_lbs)
                load = effective_load(weight, equipment, bw)
                e1rm = (
                    float(s.e1rm_lbs)
                    if s.e1rm_lbs is not None
                    else effective_e1rm(weight, s.reps, equipment, bw)
                )
                ws.append([
                    str(w.date),
                    w.name or "",
                    bw if bw else None,
                    w.sleep_quality,
                    w.fatigue_level,
                    exercise_name,
                    equipment or "",
                    s.set_number,
                    weight,
                    s.reps,
                    float(s.rpe) if s.rpe is not None else None,
                    s.set_type,
                    round(load, 1),
                    round(load * s.reps, 1),
                    round(e1rm, 1) if e1rm else None,
                    bool(s.is_pr),
                    w.notes or "",
                    we.notes or "",
                    s.notes or "",
                ])

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
